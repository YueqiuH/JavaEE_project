from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "工程文档-指导版"
SQL = ROOT / "JavaEE_project" / "database" / "baseline" / "init.sql"
DATE = "2026年7月17日"
PROJECT = "智慧校园服务平台"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False, size=9):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.name = "宋体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            r.font.size = Pt(size)
            r.bold = bold


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True, 9)
        shade(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def base_doc(title, version="V1.0"):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.0)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color in [("Title", 24, "1F4E79"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2F75B5"), ("Heading 3", 11, "404040")]:
        s = styles[name]
        s.font.name = "黑体"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run(PROJECT)
    r.bold, r.font.size, r.font.name = True, Pt(26), "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold, r.font.size, r.font.name = True, Pt(22), "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    doc.add_paragraph()
    add_table(doc, ["文档属性", "内容"], [
        ["项目名称", PROJECT], ["文档名称", title], ["版本", version],
        ["编制部门", "项目组"], ["编制日期", DATE], ["适用对象", "项目经理、开发人员、测试人员、维护人员"]
    ], [4, 12])
    doc.add_page_break()
    doc.add_heading("文档信息", 1)
    add_table(doc, ["标题", "作者", "创建日期", "最后更新", "版本"], [[f"{PROJECT}-{title}", "项目组", DATE, DATE, version]])
    doc.add_heading("修订历史", 1)
    add_table(doc, ["日期", "版本", "说明", "作者"], [[DATE, version, "依据项目源码、实体类及数据库初始化脚本编制正式版", "项目组"]])
    return doc


def add_contents(doc, items):
    doc.add_heading("目录", 1)
    for x in items:
        p = doc.add_paragraph(x)
        p.paragraph_format.left_indent = Cm(0.5 if x.count(".") <= 1 else 1.0)
    doc.add_page_break()


def parse_sql():
    text = SQL.read_text(encoding="utf-8")
    tables = []
    pattern = re.compile(r"CREATE TABLE IF NOT EXISTS `([^`]+)`\s*\((.*?)\) ENGINE=.*?COMMENT='([^']*)';", re.S)
    for name, body, comment in pattern.findall(text):
        fields = []
        pk_cols, unique_cols = set(), set()
        for line in body.splitlines():
            line = line.strip().rstrip(",")
            m = re.match(r"`([^`]+)`\s+([^\s,]+(?:\([^)]*\))?)(.*)", line)
            if m:
                col, typ, rest = m.groups()
                cm = re.search(r"COMMENT\s+'([^']*)'", rest)
                fields.append([cm.group(1) if cm else col, col, typ.upper(), "是" if "NOT NULL" in rest else "否", "", "否", "唯一" if "UNIQUE" in rest else ""])
            pm = re.search(r"PRIMARY KEY\s*\(([^)]*)\)", line)
            if pm:
                pk_cols.update(re.findall(r"`([^`]+)`", pm.group(1)))
            um = re.search(r"UNIQUE KEY.*?\(([^)]*)\)", line)
            if um:
                unique_cols.update(re.findall(r"`([^`]+)`", um.group(1)))
        for f in fields:
            f[4] = "是" if f[1] in pk_cols else "否"
            if f[1] in unique_cols:
                f[6] = "联合唯一"
            if f[1].endswith("_id") and not f[4] == "是":
                f[5] = "逻辑关联"
        tables.append((name, comment, fields))
    return tables


def database_doc():
    tables = parse_sql()
    doc = base_doc("数据库设计说明书")
    add_contents(doc, ["1 引言", "1.1 编写目的", "1.2 背景与范围", "1.3 定义与参考资料", "2 外部设计", "2.1 数据库与命名规范", "2.2 表分类", "3 结构设计", "3.1 概念结构", "3.2 物理表设计", "4 运用设计", "4.1 数据字典与完整性", "4.2 安全、备份与性能设计"])
    doc.add_heading("1 引言", 1)
    doc.add_heading("1.1 编写目的", 2)
    doc.add_paragraph("本文档定义智慧校园服务平台的数据存储结构、字段含义、约束、逻辑关系及运行管理要求，为后端编码、接口联调、测试数据准备、数据库部署和后续维护提供统一依据。")
    doc.add_heading("1.2 背景与范围", 2)
    doc.add_paragraph("系统采用前后端分离架构，后端为 Java/Spring Boot、MyBatis-Plus，数据库为 MySQL。数据库名称为 school_spring，覆盖用户与组织、教务教学、学生事务、协同办公、校园信息与统计等业务。")
    doc.add_heading("1.3 定义与参考资料", 2)
    add_table(doc, ["术语/资料", "说明"], [["CDM", "概念数据模型"], ["PDM", "物理数据模型"], ["PK/UK", "主键/唯一约束"], ["逻辑关联", "当前脚本以字段约定表达关联，未声明物理 FOREIGN KEY"], ["参考资料", "database/baseline/init.sql、campus-contract 实体类、application.yml、项目 README"]])
    doc.add_heading("2 外部设计", 1)
    doc.add_heading("2.1 数据库与命名规范", 2)
    add_table(doc, ["项目", "约定"], [["数据库", "MySQL；库名 school_spring；字符集 utf8mb4；排序规则 utf8mb4_general_ci"], ["存储引擎", "InnoDB"], ["命名", "表名与字段名使用小写 snake_case；主键通常为 表意前缀_id"], ["时间", "日期使用 DATE，时刻使用 TIME，时间戳使用 DATETIME"], ["状态", "INT/VARCHAR 枚举值；具体取值记录在字段注释中"], ["金额", "DECIMAL(10,2)，避免浮点误差"]])
    doc.add_heading("2.2 表分类", 2)
    categories = [
        ["共享基础", "user、grade、student、course、menu、department、major"],
        ["教务教学", "classroom、schedule、schedule_change、course_selection、course_capacity、score、exam、exam_room、invigilation、resit_apply、graduation_topic、graduation_selection、graduation_report"],
        ["学生事务", "student_status_change、scholarship、evaluation、competition、competition_team、competition_member、lab、lab_booking、psychology_questionnaire、psychology_warning"],
        ["协同办公", "fee、payment、asset、work_plan、document、document_approval、meeting、meeting_attendee、notification"],
        ["信息与统计", "enrollment、news、forum_post、forum_comment"]]
    add_table(doc, ["类型", "数据表"], categories, [3, 13])
    doc.add_paragraph(f"初始化脚本共定义 {len(tables)} 张业务表。")
    doc.add_heading("3 结构设计", 1)
    doc.add_heading("3.1 概念结构", 2)
    doc.add_paragraph("核心实体关系如下：院系包含专业；年级关联学生；学生围绕选课、成绩、补考、评教、实验室预约、竞赛、奖助贷、学籍异动、缴费形成业务记录；课程关联排课、容量、考试与成绩；公文关联审批记录，会议关联参会人员，帖子关联评论。各关联字段在当前基线脚本中采用逻辑外键，由业务层校验一致性。")
    doc.add_heading("3.2 物理表设计", 2)
    for idx, (name, comment, fields) in enumerate(tables, 1):
        doc.add_heading(f"3.2.{idx} {comment}（{name}）", 3)
        add_table(doc, ["字段中文名", "字段名", "类型", "非空", "主键", "外键/关联", "约束"], fields, [3.3, 3, 2.5, 1.2, 1.2, 2, 2])
    doc.add_heading("4 运用设计", 1)
    doc.add_heading("4.1 数据字典与完整性", 2)
    doc.add_paragraph("主键均用于唯一标识记录；course_selection 通过 student_id、course_id、semester 联合唯一约束防止同一学生同学期重复选课。必填字段依靠 NOT NULL 约束，业务枚举、时间范围、容量、金额及关联记录存在性由服务层进一步校验。删除被引用数据前应先检查逻辑关联，避免产生孤立记录。")
    doc.add_heading("4.2 安全、备份与性能设计", 2)
    add_table(doc, ["方面", "设计要求"], [["访问控制", "数据库账号最小权限；应用接口依据用户类型和资源归属鉴权；管理操作记录审计日志"], ["敏感数据", "密码使用强散列保存，不以明文存储；配置中的账号、口令和 API 密钥通过环境变量或密钥服务注入"], ["输入安全", "使用 MyBatis 参数绑定，禁止拼接用户输入；对分页大小、排序字段和上传内容设置白名单"], ["备份恢复", "每日增量、每周全量备份；定期进行恢复演练；发布前备份并保留可回滚脚本"], ["索引", "主键自动建索引；按查询频率为 username、student_no、semester、status、create_time 及常用关联字段建立索引"], ["连接与事务", "使用连接池；缴费、容量变更、审批等多表操作使用事务并做好并发控制"]])
    return doc


FUNCTIONS = [
    ["身份认证与导航", "账号密码登录、错误凭据处理、按用户类型展示菜单、退出登录", "1"],
    ["基础数据管理", "用户、院系、专业、年级、学生、课程的查询及维护", "1"],
    ["教务教学", "课表与调课、选课容量、考试考场、监考、成绩、补考、毕业设计", "1"],
    ["学生事务", "学籍异动、奖助贷、评教、竞赛、实验室预约、心理问卷与预警", "1"],
    ["协同办公", "费用支付、资产、工作计划、公文审批、会议通知", "1"],
    ["校园信息", "招生统计、新闻公告、论坛帖子与评论", "2"],
    ["AI 辅助", "AI 报告、学习辅助、心理分析、公文辅助审批的正常与降级处理", "2"],
]


def test_plan_doc():
    doc = base_doc("测试计划")
    add_contents(doc, ["1 简介", "1.1 目的与背景", "1.2 范围与定义", "2 测试需求", "2.1 功能测试需求", "2.2 非功能测试需求", "3 测试策略", "3.1 测试类型", "3.2 工具与环境", "3.3 准入、暂停、恢复与通过准则", "4 资源与职责", "5 测试活动与交付物", "6 风险与应对"])
    doc.add_heading("1 简介", 1)
    doc.add_heading("1.1 目的与背景", 2)
    doc.add_paragraph("本计划用于组织智慧校园服务平台的系统测试，明确测试范围、方法、资源、进度、完成标准和风险。平台服务学生、教师、教务人员和管理员，覆盖基础数据、教学管理、学生服务、协同办公、信息交流及 AI 辅助功能。")
    doc.add_heading("1.2 范围与定义", 2)
    doc.add_paragraph("本轮以当前 JavaEE_project 可部署版本为对象，执行页面、接口、数据库、权限、兼容性和基础性能测试。第三方浏览器、MySQL、Redis 和外部 AI 服务本身不在测试范围内，但验证其接口集成、超时与失败降级。优先级：1=高，2=中，3=低。缺陷等级：S1 阻断、S2 严重、S3 一般、S4 建议。")
    doc.add_heading("2 测试需求", 1)
    doc.add_heading("2.1 功能测试需求", 2)
    add_table(doc, ["测试需求", "测试需求简介", "优先级"], FUNCTIONS)
    doc.add_heading("2.2 非功能测试需求", 2)
    add_table(doc, ["类别", "条件", "指标"], [["性能", "常用列表/详情接口，正常负载", "95% 请求响应时间≤3秒，页面首次可交互≤5秒"], ["并发一致性", "选课、支付、审批等重复提交", "不产生重复记录、超卖或重复状态迁移"], ["安全", "未登录、越权、非法参数、注入型输入", "拒绝访问且不泄漏堆栈、密钥和敏感字段"], ["兼容性", "Windows 10/11；Chrome、Edge、Firefox 当前稳定版", "核心流程可完成，布局无阻断性错位"], ["可靠性", "外部 AI/Redis 暂时不可用", "核心非 AI 业务保持可用，给出明确错误或降级提示"], ["数据", "增删改查、事务回滚", "数据准确、约束有效、失败操作不遗留脏数据"]])
    doc.add_heading("3 测试策略", 1)
    doc.add_heading("3.1 测试类型", 2)
    add_table(doc, ["类型", "方法", "重点"], [["功能测试", "等价类、边界值、错误推测、场景法", "正常、空数据、非法参数、重复提交和状态流转"], ["接口测试", "按 OpenAPI/Controller 契约发送请求并核对响应与落库", "HTTP 方法、参数、统一响应、分页、异常码"], ["权限测试", "不同 user_type 与未登录身份交叉访问", "菜单可见性、接口鉴权、资源归属"], ["数据库测试", "核对约束、事务及逻辑关联", "唯一性、金额精度、容量一致性、孤立记录"], ["性能测试", "逐步加压并采集响应时间和错误率", "登录、列表查询、选课、支付等关键链路"], ["兼容性测试", "多浏览器执行冒烟集", "登录、导航、表格、表单、对话框和图表"], ["回归测试", "缺陷修复后执行相关模块及 P1 冒烟集", "防止修复引入旁路问题"]])
    doc.add_heading("3.2 工具与环境", 2)
    add_table(doc, ["用途", "工具/环境", "说明"], [["运行环境", "JDK 17+、Maven、Node.js、npm", "以项目 pom.xml 和 package.json 锁定依赖为准"], ["服务", "Spring Boot :8888、MySQL 8.x、Redis 6+", "数据库 school_spring，使用独立测试数据"], ["接口测试", "Swagger UI / Postman / curl", "保存请求、响应与截图证据"], ["性能测试", "Apache JMeter 5.6+", "逐步加压，记录 P95、吞吐量、错误率"], ["前端兼容", "Chrome、Edge、Firefox", "开发者工具检查网络请求与控制台错误"], ["缺陷管理", "项目缺陷表或 Git Issue", "记录版本、环境、步骤、证据、严重度和状态"]])
    doc.add_heading("3.3 准入、暂停、恢复与通过准则", 2)
    add_table(doc, ["阶段", "准则"], [["准入", "待测版本可构建并启动；数据库脚本可执行；核心页面可访问；测试数据和账号就绪"], ["暂停", "环境连续不可用、数据基线损坏或出现阻断后续测试的 S1 缺陷"], ["恢复", "环境与基线恢复，阻断缺陷已修复并通过冒烟验证"], ["通过", "计划用例执行率100%；P1 用例通过率100%，总通过率≥95%；S1/S2 缺陷清零；S3 修复率≥80%；性能与兼容性指标达标；文档与实现一致"]])
    doc.add_heading("4 资源与职责", 1)
    add_table(doc, ["角色", "建议人数", "职责"], [["测试负责人", "1", "维护计划、组织评审、控制范围和发布判定"], ["测试设计/执行", "2", "设计用例、准备数据、执行与复测、收集证据"], ["开发人员", "前后端各2", "定位和修复缺陷、补充单元/接口测试、支持联调"], ["项目负责人", "1", "协调资源、接受风险、批准测试结论"]])
    doc.add_heading("5 测试活动与交付物", 1)
    add_table(doc, ["活动", "工作", "时间", "交付物"], [["计划", "评审范围、需求、风险和资源", "第1天", "测试计划"], ["设计", "编写用例、数据和检查表", "第2—3天", "测试用例、测试数据"], ["执行", "冒烟、功能、接口、权限、兼容与性能测试", "第4—6天", "测试日志、缺陷单、证据"], ["回归", "复测缺陷并执行核心回归", "第7天", "回归记录"], ["评估", "统计覆盖率、通过率和遗留风险", "第8天", "测试总结/分析报告"]])
    doc.add_heading("6 风险与应对", 1)
    add_table(doc, ["风险", "影响", "应对"], [["部分业务模块仅有页面或实体，接口实现不完整", "端到端用例阻塞", "先完成静态检查和页面测试，记录阻塞并按模块补齐接口后回归"], ["测试账号/数据不足", "权限和状态流转覆盖不足", "准备四类用户及可重复恢复的数据基线"], ["外部 AI 服务波动", "AI 用例不稳定", "使用超时、错误和模拟响应验证降级；与核心发布准则分离"], ["配置含固定凭据", "安全与环境迁移风险", "测试环境使用环境变量和专用低权限凭据，发布前执行秘密扫描"], ["逻辑外键缺少数据库级约束", "可能产生孤立数据", "在服务层校验并增加数据库一致性巡检用例"]])
    return doc


CASES = [
    ("TC-AUTH-001", "账号密码登录成功", "P1", "存在有效学生账号，打开登录页", "输入正确用户名和密码；点击登录", "返回成功并进入主界面；显示该身份可见菜单"),
    ("TC-AUTH-002", "错误密码登录失败", "P1", "存在有效账号", "输入正确用户名和错误密码；点击登录", "拒绝登录并给出明确提示；不生成有效会话"),
    ("TC-AUTH-003", "空用户名/密码校验", "P1", "打开登录页", "保持用户名或密码为空；点击登录", "前端或后端提示必填，不提交无效登录"),
    ("TC-AUTH-004", "按身份展示菜单", "P1", "准备学生、教师、教务、管理员账号", "分别登录并展开导航菜单", "仅显示 user_type 对应菜单；无权限菜单不可访问"),
    ("TC-AUTH-005", "退出登录", "P1", "用户已登录", "点击退出；再次访问受保护页面", "会话被清除并返回登录页，受保护请求被拒绝"),
    ("TC-BASE-001", "院系专业查询", "P1", "数据库存在院系与专业", "进入院系专业页；按名称查询；查看专业", "列表、筛选结果及所属院系关系正确"),
    ("TC-BASE-002", "新增专业必填校验", "P2", "管理员已登录", "新增专业但不选择院系或不填名称；保存", "阻止保存并定位必填字段"),
    ("TC-BASE-003", "学生列表空数据", "P2", "使用无匹配结果的筛选条件", "执行学生查询", "显示空状态而非报错，分页总数为0"),
    ("TC-TEACH-001", "课表查询", "P1", "存在指定学期排课", "选择学期和星期；查询课表", "显示课程、教师、教室、节次和周次，数据与库一致"),
    ("TC-TEACH-002", "调课申请状态流转", "P1", "教师已登录且存在排课", "提交调课；教务审批通过", "状态由待审批变为已通过，审批人和意见被记录"),
    ("TC-TEACH-003", "非法节次校验", "P1", "进入排课或调课表单", "设置开始节次大于结束节次；提交", "拒绝提交并提示节次范围错误"),
    ("TC-TEACH-004", "选课成功", "P1", "学生已登录；课程有余量且未选", "选择课程并确认", "生成已选记录，current_count 加1且不超过 max_capacity"),
    ("TC-TEACH-005", "重复选课", "P1", "学生本学期已选该课程", "再次提交相同课程选课", "请求失败；联合唯一约束有效；容量不重复增加"),
    ("TC-TEACH-006", "满容量选课", "P1", "课程 current_count=max_capacity", "学生提交选课", "提示容量已满，不新增记录"),
    ("TC-TEACH-007", "退课", "P1", "存在状态为已选的记录", "执行退课并刷新", "状态变为退选，容量计数一致且不可重复退课"),
    ("TC-TEACH-008", "考试安排查询", "P1", "存在考试、考场及监考数据", "按学期/课程查询", "正确显示考试时间、考场、容量和监考教师"),
    ("TC-TEACH-009", "成绩边界值", "P1", "教师进入成绩录入", "分别录入0、100、-1、101", "0和100可保存；-1和101被拒绝"),
    ("TC-TEACH-010", "毕业设计流程", "P2", "教师发布课题，学生可选题", "学生选题；教师审核；提交报告", "选题与报告状态按规则流转并保留审核信息"),
    ("TC-STU-001", "奖助贷申请", "P1", "学生已登录", "填写类型、标题、理由并提交", "生成待审核申请，申请人和时间正确"),
    ("TC-STU-002", "评教评分边界", "P1", "学生有可评价课程", "输入1、5及越界0、6", "1和5有效；越界值被拒绝；评价关联学生、教师和课程"),
    ("TC-STU-003", "竞赛组队报名", "P2", "竞赛报名中且未达队伍上限", "创建队伍；添加成员；提交审核", "队伍为待审核，队长和成员关系正确"),
    ("TC-STU-004", "实验室预约", "P1", "实验室开放", "选择日期及连续节次预约", "生成预约记录；时间范围正确；页面反馈成功"),
    ("TC-STU-005", "心理问卷与预警", "P1", "学生已登录", "提交高风险答案", "保存答案与分析；按规则产生预警且仅授权人员可见"),
    ("TC-OFFICE-001", "费用账单查询", "P1", "学生存在未支付账单", "进入缴费页；按学期/状态筛选", "金额、类型、到期日和状态准确"),
    ("TC-OFFICE-002", "支付成功与重复提交", "P1", "存在未支付账单", "提交支付；立即重复提交", "仅产生一条有效支付记录，账单变为已支付，金额一致"),
    ("TC-OFFICE-003", "资产申请审批", "P1", "资产在库，申请人和审批人已登录", "提交领用申请；审批通过", "审批状态更新，资产状态/领用人正确"),
    ("TC-OFFICE-004", "工作计划维护", "P2", "教职工已登录", "新增周计划；编辑内容；标记完成", "内容和日期保存，状态正确变更"),
    ("TC-OFFICE-005", "公文审批", "P1", "存在审批中公文", "审批人填写意见并同意", "新增审批记录，公文状态/下一审批人按审批链更新"),
    ("TC-OFFICE-006", "会议通知与回复", "P1", "发起人创建会议并选择参会人", "发布会议；参会人回复参会", "会议、参会记录和通知生成；回复时间及状态正确"),
    ("TC-INFO-001", "新闻发布与置顶", "P2", "管理员已登录", "新增新闻并设为置顶；刷新列表", "新闻保存并优先显示，发布人和时间正确"),
    ("TC-INFO-002", "论坛发帖评论", "P2", "用户已登录", "发布帖子；另一用户评论", "帖子和评论正确关联，计数/状态展示正确"),
    ("TC-AI-001", "AI 服务异常降级", "P2", "模拟 AI 超时或返回错误", "提交 AI 分析/审批请求", "页面给出可理解提示，不泄漏密钥或堆栈；普通业务仍可用"),
    ("TC-SEC-001", "未登录接口访问", "P1", "清除登录状态", "直接请求受保护的新增、修改、删除接口", "返回未认证响应，数据库无变化"),
    ("TC-SEC-002", "越权访问", "P1", "以学生身份登录", "请求管理员资产审批或基础数据维护接口", "返回无权限响应且记录必要审计信息"),
    ("TC-SEC-003", "注入型输入", "P1", "打开查询和文本表单", "输入引号、SQL 片段及超长文本", "参数化处理，不执行注入；超长输入被校验；响应不泄漏 SQL"),
    ("TC-PERF-001", "列表查询性能", "P2", "准备不少于1万条测试记录", "50并发持续5分钟查询常用列表", "错误率<1%，P95≤3秒，无连接池耗尽"),
    ("TC-COMP-001", "浏览器兼容冒烟", "P2", "Windows 10/11 与三种浏览器", "执行登录、导航、查询、新增、对话框和图表流程", "核心流程一致可用，无阻断性布局或脚本错误"),
]


def test_cases_doc():
    doc = base_doc("测试用例")
    add_contents(doc, ["1 测试环境", "2 测试数据与执行规则", "3 功能与非功能测试用例", "4 需求覆盖矩阵"])
    doc.add_heading("1 测试环境", 1)
    add_table(doc, ["项目", "配置"], [["客户端", "Windows 10/11；Chrome、Edge、Firefox 当前稳定版；分辨率 1366×768 及以上"], ["前端", "Vue 3 + Vite，按 package-lock.json 安装依赖"], ["后端", "JDK 17+、Spring Boot，服务端口 8888"], ["数据与缓存", "MySQL 8.x / school_spring / utf8mb4；Redis 6+"], ["网络", "本机或测试局域网；AI 异常用例允许模拟超时/断网"], ["证据", "记录版本、执行人、时间、实际结果、通过/失败、截图或日志位置"]])
    doc.add_heading("2 测试数据与执行规则", 1)
    doc.add_paragraph("准备学生、教师、教务、管理员四类账号；准备有余量/满容量课程、未支付/已支付账单、待审批公文与资产、开放实验室及不同风险问卷数据。每个破坏性用例执行前恢复基线或使用独立数据。标记为 P1 的用例组成发布冒烟集。实际执行结果不得预填，执行后由测试人员填写。")
    doc.add_heading("3 功能与非功能测试用例", 1)
    for i, (cid, title, pri, pre, steps, expected) in enumerate(CASES, 1):
        doc.add_heading(f"3.{i} {title}", 2)
        add_table(doc, ["字段", "内容"], [["测试用例编号", cid], ["版本/优先级", f"V1.0 / {pri}"], ["测试环境", "标准测试环境"], ["前置条件", pre], ["测试步骤", steps.replace("；", "\n")], ["预期结果", expected.replace("；", "\n")], ["实际结果", "执行时填写"], ["测试结论", "□通过  □失败  □阻塞"], ["备注/证据", "执行时填写"]], [3.5, 12.5])
    doc.add_heading("4 需求覆盖矩阵", 1)
    groups = {}
    for c in CASES:
        prefix = c[0].split("-")[1]
        groups.setdefault(prefix, []).append(c[0])
    names = {"AUTH":"身份认证与导航", "BASE":"基础数据", "TEACH":"教务教学", "STU":"学生事务", "OFFICE":"协同办公", "INFO":"校园信息", "AI":"AI辅助", "SEC":"安全与权限", "PERF":"性能", "COMP":"兼容性"}
    add_table(doc, ["测试需求", "覆盖用例", "覆盖结论"], [[names[k], "、".join(v), "已设计"] for k, v in groups.items()])
    return doc


def save(doc, name):
    path = OUT / name
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [save(database_doc(), "4.数据库设计说明书.docx"), save(test_plan_doc(), "6.测试计划.docx"), save(test_cases_doc(), "7.测试用例.docx")]
    for p in paths:
        print(p)
